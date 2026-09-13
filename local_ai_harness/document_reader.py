from __future__ import annotations

from pathlib import Path


class DocumentReadError(ValueError):
    pass


class DocumentReader:
    def read(self, path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            return self._read_pdf(path)
        if suffix == ".docx":
            return self._read_docx(path)
        if suffix == ".doc":
            raise DocumentReadError("Legacy .doc files are not supported in the MVP; convert to .docx")
        return self._read_text(path)

    def _read_text(self, path: Path) -> str:
        return path.read_text(encoding="utf-8", errors="replace")

    def _read_pdf(self, path: Path) -> str:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise DocumentReadError("PDF support requires pypdf; install requirements.txt") from exc

        try:
            reader = PdfReader(str(path))
            pages: list[str] = []
            for index, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                pages.append(f"--- PDF PAGE {index} ---\n{text.strip()}")
            return "\n\n".join(pages).strip()
        except Exception as exc:
            raise DocumentReadError(f"Failed to read PDF: {exc}") from exc

    def _read_docx(self, path: Path) -> str:
        try:
            from docx import Document
        except ImportError as exc:
            raise DocumentReadError("Word support requires python-docx; install requirements.txt") from exc

        try:
            doc = Document(str(path))
            parts: list[str] = []
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if text:
                    parts.append(text)
            for table_index, table in enumerate(doc.tables, start=1):
                parts.append(f"--- DOCX TABLE {table_index} ---")
                for row in table.rows:
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    parts.append(" | ".join(cells))
            return "\n".join(parts).strip()
        except Exception as exc:
            raise DocumentReadError(f"Failed to read Word document: {exc}") from exc

